"""Combining nexus batteries over start weeks (PLAN.md §4 D317) — the pure half.

`scripts/nexus/nexus_combine.py` reads published batteries and averages them; this
pins what it does with them: only one baseline combines, the score is the mean
of the start weeks, the interval is never narrower than any correlation allows,
and the rank range separates a clear leader from a tie inside the noise.
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("nexus_combine", ROOT / "scripts" / "nexus" / "nexus_combine.py")
nc = importlib.util.module_from_spec(spec)
spec.loader.exec_module(nc)


def battery(week: int, values: dict[tuple[str, str], tuple[float, float]], **over) -> dict:
    b = {"run_id": f"run-{week}", "start_week": week, "requested_start_week": week,
         "model_validation_id": "m", "policy_version_id": "p", "dataset_version_id": "d",
         "hash_inputs": "h", "code_version": "scsim-0.8.3+x", "seed": 1, "replications": 10,
         "durations": [4], "warmup_week": 10, "window_weeks": 52, "finished_at": "t",
         "clean_lost_sales_eur": 100.0, "model_code": "Q", "model_name": "M",
         "dataset_version_no": 4, "policy_version_no": 2,
         "rows": [{"entity_type": k, "entity_id": i,
                   "metrics": {"average_eur": m, "by_duration": {"4": {"mean": m, "ci_halfwidth": hw, "n": 10}}}}
                  for (k, i), (m, hw) in values.items()]}
    b.update(over)
    return b


def test_only_the_start_week_may_differ():
    a, b = battery(14, {}), battery(26, {}, code_version="scsim-0.8.2+y")
    with pytest.raises(SystemExit, match="code_version"):
        nc.check_same_baseline([a, b])
    with pytest.raises(SystemExit, match="same start week"):
        nc.check_same_baseline([battery(26, {}), battery(26, {}, run_id="other")])
    nc.check_same_baseline([battery(14, {}), battery(26, {})])


def test_score_is_the_mean_and_the_interval_the_widest_possible():
    v14 = {("supplier", "A"): (100.0, 20.0)}
    v26 = {("supplier", "A"): (300.0, 40.0)}
    row = nc.combine([battery(14, v14), battery(26, v26)])["supplier"][0]
    assert row["combined_eur"] == 200.0
    # perfect correlation: hw of the mean = mean of the hws (30), never sqrt-shrunk (≈22.4)
    assert row["ci_halfwidth_eur"] == 30.0
    assert row["by_week"][14] == {"eur": 100.0, "ci_halfwidth_eur": 20.0, "rank": 1}


def test_rank_range_separates_a_leader_from_a_tie():
    vals = {("supplier", "LEAD"): (1000.0, 10.0), ("supplier", "T1"): (500.0, 200.0),
            ("supplier", "T2"): (495.0, 200.0), ("supplier", "ZERO"): (0.0, 0.0)}
    rows = {r["id"]: r for r in nc.combine([battery(14, vals), battery(26, vals)], top_n=1)["supplier"]}
    assert (rows["LEAD"]["rank_low"], rows["LEAD"]["rank_high"], rows["LEAD"]["p_top"]) == (1, 1, 1.0)
    assert nc.reliability(rows["LEAD"], top_n=1) == "solid top"
    # T1 and T2 swap places within the noise: each has a range, not a fixed rank
    assert rows["T1"]["rank_low"] == 2 and rows["T1"]["rank_high"] >= 3
    assert rows["T2"]["rank_low"] == 2 and rows["T2"]["rank_high"] >= 3
    assert rows["ZERO"]["no_loss"] and rows["ZERO"]["rank_low"] is None
    assert nc.reliability(rows["ZERO"]) == "no loss"


def test_timing_dependence_is_visible_per_start_week():
    early = {("material", "X"): (900.0, 1.0), ("material", "Y"): (100.0, 1.0)}
    late = {("material", "X"): (100.0, 1.0), ("material", "Y"): (800.0, 1.0)}
    rows = {r["id"]: r for r in nc.combine([battery(14, early), battery(50, late)])["material"]}
    assert rows["X"]["rank"] == 1 and rows["X"]["by_week"][50]["rank"] == 2
    assert rows["Y"]["by_week"][50]["rank"] == 1


def test_combination_is_reproducible_and_refuses_missing_targets():
    vals = {("supplier", "A"): (10.0, 5.0), ("supplier", "B"): (9.0, 5.0)}
    bs = [battery(14, vals), battery(26, vals)]
    assert nc.combine(bs) == nc.combine(bs)
    short = battery(38, {("supplier", "A"): (10.0, 5.0)})
    with pytest.raises(SystemExit, match="missing"):
        nc.combine(bs + [short])


def test_the_workbook_has_one_sheet_per_kind_and_its_provenance(tmp_path):
    vals = {("supplier", "A"): (10.0, 5.0), ("material", "M"): (7.0, 1.0)}
    bs = [battery(14, vals), battery(26, vals)]
    files = nc.write_outputs(bs, nc.combine(bs), "Project X", tmp_path)
    from openpyxl import load_workbook
    wb = load_workbook(files["xlsx"])
    assert wb.sheetnames == ["Suppliers", "Materials", "How it was made"]
    assert wb["Suppliers"]["B2"].value == "A"
    assert any("run-26" in str(c.value) for c in wb["How it was made"]["B"])
